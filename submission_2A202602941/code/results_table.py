"""Lưu JSON và xuất bảng thí nghiệm bằng template gốc.

Nhiệm vụ: lưu kết quả từng lần chạy ra JSON, rồi điền vào experiments.xlsx từ mẫu
templates/experiment_table_template.xlsx (đừng gõ tay hàng chục dòng, rất dễ sai).

Tên cột của sheet "Experiments" (giữ nguyên, đúng thứ tự mẫu):
    exp_id, group, description, loss, optimizer, lr, weight_decay, batch, epochs, hidden, dropout,
    clip_norm, precision, init, seed, step0_loss, best_val_loss, best_epoch, final_train_loss,
    final_val_loss, val_acc, val_macro_f1, time_per_epoch_s, peak_mem_MB, diverged,
    eval_acc, eval_macro_f1, figure_file, notes
(các cột công thức ở cuối bảng mẫu tự tính, đừng ghi đè)
"""
from __future__ import annotations

import json
from pathlib import Path
from copy import copy
import math
import numpy as np
import openpyxl
from openpyxl.formula.translate import Translator
from openpyxl.workbook.properties import CalcProperties

FORMULAS = {"step0_gap_vs_lnC", "gap_val_minus_train", "delta_val_f1_vs_base", "beyond_noise"}


def _safe(value):
    if isinstance(value, dict):
        return {str(k): _safe(v) for k, v in value.items()}
    if isinstance(value, (list, tuple, np.ndarray)):
        return [_safe(v) for v in value]
    if isinstance(value, np.generic):
        return _safe(value.item())
    if isinstance(value, float) and not math.isfinite(value):
        return None
    return value


def save_result(result: dict, results_dir: str = "../results") -> str:
    """Ghi result["cfg"], result["history"], result["summary"] (KHÔNG ghi best_state) ra
    <results_dir>/<exp_id>.json. Trả về đường dẫn file. Tạo thư mục nếu chưa có."""
    exp_id = result["cfg"]["exp_id"]
    if not isinstance(exp_id, str) or not exp_id or any(c in exp_id for c in "/\\") or exp_id in (".", ".."):
        raise ValueError("exp_id must be a single file name")
    path = Path(results_dir) / f"{exp_id}.json"
    path.parent.mkdir(parents=True, exist_ok=True)
    payload = {k: _safe(result[k]) for k in ("cfg", "history", "summary")}
    temporary = path.with_suffix(".json.tmp")
    temporary.write_text(json.dumps(payload, ensure_ascii=False, indent=2, allow_nan=False) + "\n", encoding="utf-8")
    temporary.replace(path)
    return str(path)


def load_results(results_dir: str = "../results") -> list[dict]:
    """Đọc mọi file *.json trong results_dir, trả về danh sách dict (sắp theo exp_id)."""
    results = []
    for path in sorted(Path(results_dir).glob("*.json")):
        item = json.loads(path.read_text(encoding="utf-8"))
        if all(k in item for k in ("cfg", "history", "summary")):
            results.append(item)
    ids = [r["cfg"]["exp_id"] for r in results]
    if len(set(ids)) != len(ids):
        raise ValueError("Duplicate exp_id in result files")
    return sorted(results, key=lambda r: r["cfg"]["exp_id"])


def to_row(result: dict, eval_scores: dict | None = None, notes: str = "") -> dict:
    """Biến một kết quả thành một dòng của bảng: gộp cfg + summary (+ eval_acc, eval_macro_f1 nếu có)
    + figure_file = f"figures/{exp_id}.png". Khoá phải trùng tên cột ở đầu file.
    Chỉ truyền eval_scores cho baseline và cấu hình cuối cùng."""
    cfg, summary = result["cfg"], result["summary"]
    row = {**cfg, **summary}
    row["hidden"] = "-".join(map(str, cfg["hidden"]))
    row["clip_norm"] = "none" if cfg["clip_norm"] is None else cfg["clip_norm"]
    row["figure_file"] = f"figures/{cfg['exp_id']}.png"
    details = [notes, cfg.get("notes", ""), summary.get("divergence_reason", "") or ""]
    for key in ("momentum", "betas", "eps", "scheduler", "train_eval_samples", "activation_std", "prediction", "observation", "smoke"):
        if key in cfg:
            details.append(f"{key}={cfg[key]}")
    row["notes"] = "; ".join(str(s) for s in details if s)
    row["eval_acc"] = eval_scores["accuracy"] if eval_scores is not None else None
    row["eval_macro_f1"] = eval_scores["macro_f1"] if eval_scores is not None else None
    return _safe(row)


def write_xlsx(rows: list[dict], template_path: str, out_path: str) -> None:
    """Điền các dòng vào sheet "Experiments" của mẫu, từ dòng 2 trở xuống, rồi lưu thành out_path.

    Các bước (openpyxl):
      1. wb = openpyxl.load_workbook(template_path)   # KHÔNG dùng data_only=True (sẽ mất công thức)
      2. ws = wb["Experiments"]; đọc tiêu đề dòng 1 để biết cột nào ứng với khoá nào
      3. với mỗi row: ghi giá trị vào đúng cột; BỎ QUA các cột công thức (step0_gap_vs_lnC, gap_val_minus_train,
         delta_val_f1_vs_base, beyond_noise)
      4. wb.save(out_path)
    Sau khi lưu, mở file bằng Excel/LibreOffice để các công thức tính lại.
    """
    ids = [r["exp_id"] for r in rows]
    if len(ids) != len(set(ids)):
        raise ValueError("Each experiment must have a unique exp_id")
    wb = openpyxl.load_workbook(template_path)
    ws = wb["Experiments"]
    headers = [c.value for c in ws[1]]
    formulas = {i: ws.cell(2, i).value for i, h in enumerate(headers, 1) if h in FORMULAS}
    original_max = ws.max_row
    for cells in ws.iter_rows(min_row=2):
        for cell in cells:
            if headers[cell.column - 1] not in FORMULAS:
                cell.value = None
    for index, row in enumerate(rows, 2):
        for col, header in enumerate(headers, 1):
            cell = ws.cell(index, col)
            if index > original_max:
                cell._style = copy(ws.cell(2, col)._style)
            if header in FORMULAS:
                cell.value = Translator(formulas[col], origin=ws.cell(2, col).coordinate).translate_formula(cell.coordinate)
            else:
                value = row.get(header)
                cell.value = json.dumps(value, ensure_ascii=False) if isinstance(value, (dict, list)) else value
    last = max(61, len(rows) + 1)
    if last > 61:
        for sheet in wb:
            for cells in sheet:
                for cell in cells:
                    if cell.data_type == "f":
                        cell.value = cell.value.replace("$61", f"${last}")
    seeds = wb["Seeds"]
    baseline = [r for r in rows if r.get("group") == "baseline"]
    if len(baseline) > 5:
        raise ValueError("Template Seeds has five baseline slots; use at most five seeds")
    for index in range(2, 7):
        seeds.cell(index, 1).value = baseline[index - 2]["exp_id"] if index - 2 < len(baseline) else None
    summary_sheet = wb["Summary"]
    for index in range(2, summary_sheet.max_row + 1):
        group = summary_sheet.cell(index, 1).value
        measured = [r for r in rows if r.get("group") == group and r.get("val_macro_f1") is not None]
        if measured:
            best = max(measured, key=lambda r: r["val_macro_f1"])
            summary_sheet.cell(index, 8, f"Best val macro-F1: {best['exp_id']} = {best['val_macro_f1']:.6f}; compare seed noise before concluding.")
    wb.calculation = CalcProperties(calcId=0, fullCalcOnLoad=True, forceFullCalc=True)
    target = Path(out_path)
    target.parent.mkdir(parents=True, exist_ok=True)
    wb.save(target)
